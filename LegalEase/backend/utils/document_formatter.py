from io import BytesIO
from pathlib import Path
import html
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


ROOT = Path(__file__).resolve().parents[2]
LOGO_PATH = ROOT / "assets" / "logo.png"


def sanitize_text(text: str) -> str:
    """Clean generated text before exporting."""
    if text is None:
        return ""

    text = str(text)

    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Remove markdown code fences if Gemini accidentally returns them
    text = text.replace("```text", "")
    text = text.replace("```markdown", "")
    text = text.replace("```", "")

    return text.strip()


def _paragraphs(content: str):
    """Return non-empty document paragraphs."""
    content = sanitize_text(content)

    return [
        line.strip()
        for line in content.split("\n")
        if line.strip()
    ]


def _is_heading(text: str) -> bool:
    """Detect common legal-document headings."""
    text = text.strip()

    if not text:
        return False

    if len(text) > 100:
        return False

    if text.isupper():
        return True

    heading_words = [
        "INTRODUCTION",
        "PARTIES",
        "DEFINITIONS",
        "PURPOSE",
        "TERM",
        "PAYMENT",
        "OBLIGATIONS",
        "RESPONSIBILITIES",
        "CONFIDENTIALITY",
        "TERMINATION",
        "GOVERNING LAW",
        "DISPUTE RESOLUTION",
        "SIGNATURES",
        "SIGNATURE",
        "AGREEMENT",
    ]

    upper = text.upper()

    return any(
        upper.startswith(word)
        for word in heading_words
    )


def format_txt(
    content: str,
    document_type: str = "",
    effective_date: str = "",
) -> bytes:
    """Create a plain-text legal document."""

    content = sanitize_text(content)

    header = document_type.strip()

    if effective_date.strip():
        header += f"\nEffective Date: {effective_date.strip()}"

    if header:
        output = f"{header}\n\n{content}"
    else:
        output = content

    return output.encode("utf-8")


def format_html_preview(
    content: str,
    document_type: str = "",
    effective_date: str = "",
) -> str:
    """Create HTML suitable for a document preview."""

    content = sanitize_text(content)

    title = html.escape(document_type.strip())

    date_html = ""

    if effective_date.strip():
        date_html = (
            f'<div class="effective-date">'
            f'Effective Date: {html.escape(effective_date.strip())}'
            f"</div>"
        )

    paragraphs = []

    for paragraph in _paragraphs(content):
        escaped = html.escape(paragraph)

        if _is_heading(paragraph):
            paragraphs.append(
                f'<h3>{escaped}</h3>'
            )
        else:
            paragraphs.append(
                f"<p>{escaped}</p>"
            )

    return f"""
    <div class="legal-document">
        <h1>{title}</h1>
        {date_html}
        {''.join(paragraphs)}
    </div>
    """


def _set_cell_text(cell, text: str):
    """Set table cell text."""
    cell.text = ""

    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))

    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def format_docx(
    content: str,
    document_type: str = "",
    terms: str = "",
    effective_date: str = "",
) -> bytes:
    """
    Create a professional DOCX document.
    """

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Default font
    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # Logo
    if LOGO_PATH.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()

        try:
            run.add_picture(
                str(LOGO_PATH),
                width=Inches(1.2),
            )
        except Exception:
            pass

    # Title
    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        document_type.strip() or "LEGAL DOCUMENT"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    # Effective date
    if effective_date.strip():
        date_paragraph = document.add_paragraph()

        date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = date_paragraph.add_run(
            f"Effective Date: {effective_date.strip()}"
        )

        run.font.name = "Times New Roman"
        run.font.size = Pt(10)

    # Document content
    for line in _paragraphs(content):

        paragraph = document.add_paragraph()

        if _is_heading(line):
            run = paragraph.add_run(line)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        else:
            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    # Key terms table
    if terms.strip():

        document.add_paragraph()

        heading = document.add_paragraph()

        run = heading.add_run("KEY TERMS")
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        _set_cell_text(
            table.rows[0].cells[0],
            "Term",
        )

        _set_cell_text(
            table.rows[0].cells[1],
            "Details",
        )

        for item in terms.split(";"):

            item = item.strip()

            if not item:
                continue

            if ":" in item:
                key, value = item.split(
                    ":",
                    1,
                )

            else:
                key = item
                value = ""

            cells = table.add_row().cells

            _set_cell_text(
                cells[0],
                key.strip(),
            )

            _set_cell_text(
                cells[1],
                value.strip(),
            )

    # Footer
    footer = section.footer

    paragraph = footer.paragraphs[0]

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "Generated by LegalEase — AI-assisted legal document draft"
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(8)

    # Save to memory
    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()


class LegalEasePDF(FPDF):

    def footer(self):
        self.set_y(-15)

        self.set_font(
            "Helvetica",
            size=8,
        )

        self.cell(
            0,
            10,
            "Generated by LegalEase - AI-assisted legal document draft",
            align="C",
        )


def format_pdf(
    content: str,
    document_type: str = "",
    terms: str = "",
    effective_date: str = "",
) -> bytes:
    """Create a PDF legal document."""

    pdf = LegalEasePDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.add_page()

    # Logo
    if LOGO_PATH.exists():

        try:
            pdf.image(
                str(LOGO_PATH),
                x=85,
                y=10,
                w=40,
            )

            pdf.ln(30)

        except Exception:
            pass

    # Title
    pdf.set_font(
        "Helvetica",
        style="B",
        size=16,
    )

    pdf.multi_cell(
        0,
        10,
        document_type.strip() or "LEGAL DOCUMENT",
        align="C",
    )

    if effective_date.strip():

        pdf.set_font(
            "Helvetica",
            size=9,
        )

        pdf.multi_cell(
            0,
            7,
            f"Effective Date: {effective_date.strip()}",
            align="C",
        )

    pdf.ln(5)

    # Main content
    for line in _paragraphs(content):

        if _is_heading(line):

            pdf.set_font(
                "Helvetica",
                style="B",
                size=11,
            )

            pdf.multi_cell(
                0,
                7,
                line,
            )

        else:

            pdf.set_font(
                "Helvetica",
                size=10,
            )

            pdf.multi_cell(
                0,
                6,
                line,
            )

        pdf.ln(1)

    # Terms
    if terms.strip():

        pdf.ln(3)

        pdf.set_font(
            "Helvetica",
            style="B",
            size=11,
        )

        pdf.multi_cell(
            0,
            7,
            "KEY TERMS",
        )

        pdf.set_font(
            "Helvetica",
            size=10,
        )

        for item in terms.split(";"):

            item = item.strip()

            if item:

                pdf.multi_cell(
                    0,
                    6,
                    f"- {item}",
                )

    output = pdf.output()

    return bytes(output)